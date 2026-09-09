import io, json
from datetime import date

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, Sum, Max, Min
from django.http import HttpResponse
from django.utils import timezone
from django.template.loader import render_to_string

from pacientes.models import Paciente
from citas.models import Cita
from medicos.models import Medico
from historias.models import HistoriaClinica, Evolucion
from usuarios.models import ActivityLog
from facturacion.models import Factura
from core.decorators import permiso_requerido


@login_required
@permiso_requerido('ver_reportes')
def reportes_dashboard(request):
    return render(request, 'reportes/dashboard.html', {
        'title': 'Reportes'
    })


@login_required
@permiso_requerido('ver_reporte_pacientes')
def reporte_pacientes(request):
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    filtro = Q(activo=True)

    if desde:
        filtro &= Q(fecha_registro__gte=desde)
    if hasta:
        filtro &= Q(fecha_registro__lte=hasta)

    qs = Paciente.objects.filter(filtro)

    total = qs.count()
    hombres = qs.filter(genero='M').count()
    mujeres = qs.filter(genero='F').count()
    con_foto = qs.exclude(foto='').count()

    # Por edad
    hoy = timezone.localdate()
    menores = sum(1 for p in qs if p.edad and p.edad < 18)
    adultos = sum(1 for p in qs if p.edad and 18 <= p.edad <= 60)
    mayores = sum(1 for p in qs if p.edad and p.edad > 60)

    # Últimos pacientes
    ultimos = qs.order_by('-fecha_registro')[:10]

    return render(request, 'reportes/pacientes.html', {
        'title': 'Reporte de Pacientes',
        'total': total, 'hombres': hombres, 'mujeres': mujeres,
        'con_foto': con_foto, 'menores': menores, 'adultos': adultos,
        'mayores': mayores, 'ultimos': ultimos,
        'desde': desde, 'hasta': hasta,
    })


@login_required
@permiso_requerido('ver_reporte_citas')
def reporte_citas(request):
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    filtro = Q()

    if desde:
        filtro &= Q(fecha__gte=desde)
    if hasta:
        filtro &= Q(fecha__lte=hasta)

    qs = Cita.objects.filter(filtro)

    total = qs.count()
    pendientes = qs.filter(estado='PENDIENTE').count()
    confirmadas = qs.filter(estado='CONFIRMADA').count()
    completadas = qs.filter(estado='COMPLETADA').count()
    canceladas = qs.filter(estado='CANCELADA').count()

    # Por médico
    por_medico = Medico.objects.filter(activo=True).annotate(
        total_citas=Count('citas')
    ).order_by('-total_citas')

    # Actualizar conteo según filtro de fechas
    if desde or hasta:
        for m in por_medico:
            cs = Cita.objects.filter(doctor=m)
            if desde:
                cs = cs.filter(fecha__gte=desde)
            if hasta:
                cs = cs.filter(fecha__lte=hasta)
            m.total_citas = cs.count()
        por_medico = sorted(por_medico, key=lambda x: x.total_citas, reverse=True)

    # Próximas citas
    hoy = timezone.localdate()
    proximas = Cita.objects.filter(
        fecha__gte=hoy
    ).exclude(estado='CANCELADA').order_by('fecha', 'hora')[:10]

    return render(request, 'reportes/citas.html', {
        'title': 'Reporte de Citas',
        'total': total, 'pendientes': pendientes,
        'confirmadas': confirmadas, 'completadas': completadas,
        'canceladas': canceladas,
        'por_medico': por_medico,
        'proximas': proximas,
        'desde': desde, 'hasta': hasta,
    })


@login_required
@permiso_requerido('ver_reporte_medicos')
def reporte_medicos(request):
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    filtro_citas = Q()
    if desde:
        filtro_citas &= Q(fecha__gte=desde)
    if hasta:
        filtro_citas &= Q(fecha__lte=hasta)

    if desde or hasta:
        citas_filtradas = Cita.objects.filter(filtro_citas)
        medicos_list = []
        for m in Medico.objects.filter(activo=True):
            cs = citas_filtradas.filter(doctor=m)
            medicos_list.append({
                'medico': m,
                'total_citas': cs.count(),
                'citas_completadas': cs.filter(estado='COMPLETADA').count(),
                'citas_pendientes': cs.filter(estado='PENDIENTE').count(),
            })
        medicos_list.sort(key=lambda x: x['total_citas'], reverse=True)
        medicos = medicos_list
    else:
        medicos_qs = Medico.objects.filter(activo=True).annotate(
            total_citas=Count('citas'),
        ).order_by('-total_citas')
        medicos_list = []
        for m in medicos_qs:
            medicos_list.append({
                'medico': m,
                'total_citas': m.total_citas,
                'citas_completadas': m.citas.filter(estado='COMPLETADA').count(),
                'citas_pendientes': m.citas.filter(estado='PENDIENTE').count(),
            })
        medicos = medicos_list

    all_medicos = Medico.objects.filter(activo=True)
    total_med = all_medicos.count()
    con_usuario = all_medicos.exclude(user__isnull=True).count()

    # Por especialidad
    por_especialidad = Medico.objects.filter(activo=True).values('especialidad').annotate(
        total=Count('id')
    ).order_by('-total')

    return render(request, 'reportes/medicos.html', {
        'title': 'Reporte de Médicos',
        'medicos': medicos,
        'total_med': total_med,
        'con_usuario': con_usuario,
        'por_especialidad': por_especialidad,
        'desde': desde, 'hasta': hasta,
    })


@login_required
@permiso_requerido('ver_reporte_historias')
def reporte_historias(request):
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    filtro = Q()
    if desde:
        filtro &= Q(fecha_creacion__date__gte=desde)
    if hasta:
        filtro &= Q(fecha_creacion__date__lte=hasta)

    historias = HistoriaClinica.objects.filter(filtro)
    total_historias = historias.count()
    total_pacientes = Paciente.objects.filter(activo=True).count()
    con_historia = historias.values('paciente').distinct().count()
    sin_historia = total_pacientes - con_historia

    # Condiciones médicas frecuentes
    condiciones = {
        'Alergia a drogas': historias.filter(alergia_drogas=True).count(),
        'Diabetes': historias.filter(diabetes=True).count(),
        'Hipertensión': historias.filter(hipertension=True).count(),
        'Problemas cardíacos': historias.filter(cardiacos=True).count(),
        'Fumadores': historias.filter(fumador=True).count(),
        'Consume alcohol': historias.filter(alcohol=True).count(),
        'Toma medicación': historias.filter(medicacion=True).count(),
        'Sangrado excesivo': historias.filter(sangrado=True).count(),
    }

    # Evoluciones
    total_evoluciones = Evolucion.objects.count()
    ultimas_evoluciones = Evolucion.objects.select_related(
        'historia__paciente', 'medico'
    ).order_by('-fecha')[:10]

    return render(request, 'reportes/historias.html', {
        'title': 'Reporte de Historias',
        'total_historias': total_historias,
        'con_historia': con_historia,
        'sin_historia': sin_historia,
        'condiciones': condiciones,
        'total_evoluciones': total_evoluciones,
        'ultimas_evoluciones': ultimas_evoluciones,
        'desde': desde, 'hasta': hasta,
    })


@login_required
@permiso_requerido('ver_reporte_actividad')
def reporte_actividad(request):
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    filtro = Q()
    if desde:
        filtro &= Q(fecha__date__gte=desde)
    if hasta:
        filtro &= Q(fecha__date__lte=hasta)

    qs = ActivityLog.objects.filter(filtro)

    total_acciones = qs.count()
    creaciones = qs.filter(accion='CREAR').count()
    ediciones = qs.filter(accion='EDITAR').count()
    eliminaciones = qs.filter(accion='ELIMINAR').count()
    logins = qs.filter(accion='LOGIN').count()

    # Por modelo
    por_modelo = qs.values('modelo').annotate(
        total=Count('id')
    ).order_by('-total')

    # Por usuario
    por_usuario = qs.values('usuario__username').annotate(
        total=Count('id')
    ).order_by('-total')[:10]

    # Recientes
    recientes = qs[:20]

    return render(request, 'reportes/actividad.html', {
        'title': 'Reporte de Actividad',
        'total_acciones': total_acciones,
        'creaciones': creaciones,
        'ediciones': ediciones,
        'eliminaciones': eliminaciones,
        'logins': logins,
        'por_modelo': por_modelo,
        'por_usuario': por_usuario,
        'recientes': recientes,
        'desde': desde, 'hasta': hasta,
    })


@login_required
@permiso_requerido('ver_reporte_valores')
def reporte_valores(request):
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    periodo = request.GET.get('periodo', 'mes')  # dia, semana, mes, año

    qs = Evolucion.objects.filter(costo__isnull=False).exclude(costo=0)
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)

    total_ingresos = qs.aggregate(total=Sum('costo'))['total'] or 0
    total_consultas = qs.count()
    promedio = total_ingresos / total_consultas if total_consultas else 0
    max_costo = qs.aggregate(m=Max('costo'))['m'] or 0
    min_costo = qs.aggregate(m=Min('costo'))['m'] or 0

    # Agrupar por período
    from django.db.models.functions import TruncDay, TruncWeek, TruncMonth, TruncYear

    if periodo == 'dia':
        trunc = TruncDay('fecha')
        fmt_label = '%d/%m/%Y'
    elif periodo == 'semana':
        trunc = TruncWeek('fecha')
        fmt_label = 'Semana %V - %Y'
    elif periodo == 'ano':
        trunc = TruncYear('fecha')
        fmt_label = '%Y'
    else:  # mes
        trunc = TruncMonth('fecha')
        fmt_label = '%B %Y'

    datos = (
        qs.annotate(periodo=trunc)
        .values('periodo')
        .annotate(
            total=Sum('costo'),
            cantidad=Count('id'),
        )
        .order_by('periodo')
    )

    etiquetas = []
    valores = []
    cantidades = []
    datos_periodo = []
    for d in datos:
        label = d['periodo'].strftime(fmt_label).capitalize() if d['periodo'] else 'Sin fecha'
        etiquetas.append(label)
        valores.append(float(d['total']))
        cantidades.append(d['cantidad'])
        datos_periodo.append({
            'periodo': d['periodo'],
            'cantidad': d['cantidad'],
            'total': d['total'],
            'promedio': d['total'] / d['cantidad'] if d['cantidad'] else 0,
        })

    # Por médico
    por_medico = (
        qs.values('medico__nombres', 'medico__apellidos')
        .annotate(total=Sum('costo'), cantidad=Count('id'))
        .order_by('-total')
    )

    return render(request, 'reportes/valores.html', {
        'title': 'Reporte de Valores',
        'total_ingresos': total_ingresos,
        'total_consultas': total_consultas,
        'promedio': promedio,
        'max_costo': max_costo,
        'min_costo': min_costo,
        'periodo': periodo,
        'desde': desde, 'hasta': hasta,
        'etiquetas': json.dumps(etiquetas),
        'valores': json.dumps(valores),
        'cantidades': json.dumps(cantidades),
        'datos_periodo': datos_periodo,
        'por_medico': por_medico,
    })


# ===========================================================================
#  REPORTE DE FACTURAS
# ===========================================================================

@login_required
@permiso_requerido('ver_facturas')
def reporte_facturas(request):
    """Reporte de facturas filtrado por fechas."""
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    q = Factura.objects.all().select_related('paciente', 'medico')

    if desde:
        q = q.filter(fecha_emision__gte=desde)
    if hasta:
        q = q.filter(fecha_emision__lte=hasta)

    total_facturas = q.count()
    monto_total = q.aggregate(total=Sum('total'))['total'] or 0
    pagadas = q.filter(estado='PAGADA').count()
    monto_pagado = q.filter(estado='PAGADA').aggregate(total=Sum('total'))['total'] or 0
    pendientes = q.filter(estado='PENDIENTE').count()
    anuladas = q.filter(estado='ANULADA').count()

    # Por forma de pago
    por_forma = q.values('forma_pago').annotate(
        cantidad=Count('id'), monto=Sum('total')
    ).order_by('-cantidad')

    facturas = q.order_by('-fecha_emision')

    context = {
        'title': 'Reporte de Facturas',
        'facturas': facturas,
        'total_facturas': total_facturas,
        'monto_total': monto_total,
        'pagadas': pagadas,
        'monto_pagado': monto_pagado,
        'pendientes': pendientes,
        'anuladas': anuladas,
        'por_forma': por_forma,
        'desde': desde,
        'hasta': hasta,
    }
    return render(request, 'reportes/facturas.html', context)


@login_required
@permiso_requerido('ver_facturas')
def exportar_facturas_excel(request):
    """Exporta facturas a Excel."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    q = Factura.objects.all().select_related('paciente', 'medico')
    if desde:
        q = q.filter(fecha_emision__gte=desde)
    if hasta:
        q = q.filter(fecha_emision__lte=hasta)
    q = q.order_by('-fecha_emision')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Facturas'

    # Estilos
    bold = Font(bold=True, size=12)
    header_fill = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Fecha del reporte
    ws.merge_cells('A1:G1')
    ws['A1'] = f'Reporte de Facturas - {date.today().strftime("%d/%m/%Y")}'
    ws['A1'].font = Font(bold=True, size=14)

    if desde or hasta:
        ws.merge_cells('A2:G2')
        ws['A2'] = f'Período: {desde or "—"} a {hasta or "—"}'
        ws['A2'].font = Font(italic=True, color='666666')

    # Headers
    headers = ['N° Factura', 'Paciente', 'Médico', 'Fecha Emisión', 'Total', 'Forma Pago', 'Estado']
    start_row = 4
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Datos
    for i, f in enumerate(q, start_row + 1):
        ws.cell(i, 1, f.numero).border = thin_border
        ws.cell(i, 2, str(f.paciente)).border = thin_border
        ws.cell(i, 3, str(f.medico or '')).border = thin_border
        ws.cell(i, 4, f.fecha_emision.strftime('%d/%m/%Y')).border = thin_border
        cell_total = ws.cell(i, 5, float(f.total))
        cell_total.number_format = '$#,##0.00'
        cell_total.border = thin_border
        ws.cell(i, 6, f.get_forma_pago_display()).border = thin_border
        ws.cell(i, 7, f.get_estado_display()).border = thin_border

    # Totales
    total_row = start_row + len(q) + 1
    ws.cell(total_row, 1, 'TOTALES').font = bold
    total_sum = q.aggregate(total=Sum('total'))['total'] or 0
    cell = ws.cell(total_row, 5, float(total_sum))
    cell.font = bold
    cell.number_format = '$#,##0.00'

    # Ancho columnas
    widths = [16, 30, 30, 14, 14, 20, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=facturas_{date.today()}.xlsx'
    wb.save(response)
    return response


@login_required
@permiso_requerido('ver_facturas')
def exportar_facturas_pdf(request):
    """Exporta facturas a PDF."""
    desde = request.GET.get('desde', '')
    hasta = request.GET.get('hasta', '')
    q = Factura.objects.all().select_related('paciente', 'medico')
    if desde:
        q = q.filter(fecha_emision__gte=desde)
    if hasta:
        q = q.filter(fecha_emision__lte=hasta)
    q = q.order_by('-fecha_emision')

    total_facturas = q.count()
    monto_total = q.aggregate(total=Sum('total'))['total'] or 0

    html = render_to_string('reportes/facturas_pdf.html', {
        'facturas': q,
        'total_facturas': total_facturas,
        'monto_total': monto_total,
        'desde': desde,
        'hasta': hasta,
        'hoy': date.today(),
    })

    from weasyprint import HTML
    pdf = HTML(string=html).write_pdf()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename=facturas_{date.today()}.pdf'
    return response
